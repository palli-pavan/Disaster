"""Build the revised project report with plain academic formatting and actual dashboards."""
from pathlib import Path
import ast, json, csv, textwrap
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT=Path(__file__).resolve().parent
SUB=OUT.parent
ROOT=SUB.parent/'emergency-resource-allocation'
SHOTS=SUB/'dashboards'
URL='https://github.com/palli-pavan/Disaster'
def load(n): return json.loads((OUT/n).read_text(encoding='utf-8'))
base=load('baseline_summary.json')
shock=load('water_shock_summary.json')
surge=load('population_surge_summary.json')
local=load('local_explanation.json')
inventory=load('dataset_inventory.json')
runtime=load('runtime.json')
metrics=json.loads((ROOT/'models/metrics.json').read_text(encoding='utf-8'))
rows=list(csv.DictReader((OUT/'baseline_allocations.csv').open(encoding='utf-8')))

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.75)
sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
    s=doc.styles[name]
    s.font.name='Times New Roman'
    s.font.color.rgb=RGBColor(0,0,0)
    s.font.italic=False
    for b in s.element.xpath('.//w:pBdr'): b.getparent().remove(b)
    for fonts in s.element.xpath('.//w:rFonts'):
        for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal']
normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [('Title',21),('Heading 1',16),('Heading 2',12),('Heading 3',11.5)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    s.paragraph_format.space_after=Pt(7)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Subtitle'].font.bold=True
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
style=doc.styles.add_style('Program code',1)
style.font.name='Consolas'; style.font.size=Pt(8.3)
style.paragraph_format.line_spacing=Pt(10.4)
style.paragraph_format.space_after=Pt(0)
style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
style.paragraph_format.widow_control=False
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='ReliefPlan AI Project Report'
doc.core_properties.subject='Emergency resource allocation implementation results and dashboard screens'
doc.core_properties.author=''

chapters=[]
def p(text='',style=None): return doc.add_paragraph(text,style)
def h(text,level=2): return doc.add_heading(text,level)
def page(title):
    if chapters: doc.add_page_break()
    chapters.append(title)
    h(title,1)
def num(v,d=0): return f'{float(v):,.{d}f}'
def code(text):
    for line in text.strip('\n').splitlines():
        parts=textwrap.wrap(line.expandtabs(4),width=95,replace_whitespace=False,drop_whitespace=False,break_long_words=False,break_on_hyphens=False) or ['']
        for part in parts: p(part,'Program code')
    p()
def link(para,label=URL,url=URL):
    rid=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    element=OxmlElement('w:hyperlink'); element.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'000000'); pr.append(color)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); pr.append(underline)
    run.append(pr); text=OxmlElement('w:t'); text.text=label; run.append(text); element.append(run); para._p.append(element)
def table(headers,data,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,label in enumerate(headers): t.rows[0].cells[i].text=str(label)
    for record in data:
        cells=t.add_row().cells
        for i,val in enumerate(record): cells[i].text=str(val)
    for ri,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr(); tr.append(OxmlElement('w:cantSplit'))
        if ri==0: tr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                q=OxmlElement('w:'+side); q.set(qn('w:val'),'single'); q.set(qn('w:sz'),'4'); q.set(qn('w:color'),'D9D9D9'); b.append(q)
            pr.append(b)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EEEEEE' if ri==0 else 'FFFFFF'); pr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'85'); x.set(qn('w:type'),'dxa'); mar.append(x)
            pr.append(mar)
            for paragraph in cell.paragraphs:
                paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after=Pt(0)
                paragraph.paragraph_format.line_spacing=1.0
                for run in paragraph.runs:
                    run.font.name='Times New Roman'; run.font.size=Pt(10); run.font.color.rgb=RGBColor(0,0,0)
                    run.font.bold=ri==0
    p()
    return t
def figure(filename,caption):
    q=p(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.add_run().add_picture(str(SHOTS/filename),width=Inches(6.65))
    p(caption,'Caption')

# Cover and abstract
chapters.append('Project Report')
q=p('PROJECT REPORT');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q.runs[0].bold=True;q.runs[0].font.size=Pt(14)
q=p('AI Powered Emergency Resource Allocation and Decision Support System','Title')
q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('ReliefPlan AI','Subtitle');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('16 October 2026');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('GitHub Repository');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].bold=True
q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;link(q)
h('Abstract')
p('ReliefPlan AI is a disaster response project that estimates resource needs and recommends how available supplies can be distributed across affected areas. It uses damage observations, population records, road information and warehouse stock. A trained XGBoost model predicts damage severity, and a linear programming model allocates food, water and medical kits within stock and demand limits.')
p('The application provides a four stage workflow for describing a disaster, reviewing AI analysis, inspecting a response plan and testing changes in supply. A map displays the study areas and their priority levels. The project also includes SHAP explanations, scenario comparison and a CSV decision history.')
p('The default core run estimated 12,846 affected residents across ten areas. It allocated 15,000 food rations, 20,000 litres of water and 1,927 medical kits. A 30 percent reduction in water stock lowered water coverage from 17.30 percent to 12.11 percent. All 24 existing automated tests passed. The results demonstrate the processing pipeline on synthetic data; the current application remains a prototype with several interface and allocation policy limitations.')

page('Introduction')
h('Problem Statement')
p('During a disaster, emergency teams must distribute limited supplies among communities with different levels of damage and need. Population alone does not describe the urgency of each area. Structural damage, vulnerable age groups and blocked roads also affect response planning. A useful decision support system should bring these factors together and make its recommendations understandable.')
h('Project Objectives')
for text in ['Predict damage severity from structural damage and disaster characteristics.','Estimate the number of residents needing assistance and calculate resource demand.','Assign area priorities using damage, demand, demographic and accessibility information.','Allocate available food, water and medical supplies without exceeding stock or demand.','Display recommendations on a dashboard and compare changes in supply or population.']:
    p(text,'List Bullet')
h('Scope of the Project')
p('The project uses synthetic records for a Houston study area. The core pipeline processes ten areas, A to J. The React API restricts its analysis to seven areas, A to G, and accepts population values from the interface. The two analysis paths therefore produce different results.')
p('The recommendations support a human decision maker. The application does not control vehicles, confirm delivery or send public emergency messages. Hospital and shelter data are supplied, but separate bed and shelter assignment are outside the current core pipeline.')
h('Repository')
p('The project source is available at the following address. The report describes the implementation and outputs from the supplied Disaster folder.')
link(p())

page('System Design')
p('The computation is kept in reusable Python modules. React communicates with these modules through FastAPI, while the Streamlit interface calls the Python pipeline directly. CSV files hold the input data and decision history; trained models are loaded from the models directory.')
h('Processing Flow')
code('Input CSV files and saved model\n    -> Validate and merge data\n    -> Predict damage severity\n    -> Estimate affected population and demand\n    -> Calculate priority scores\n    -> Solve resource allocation\n    -> Generate actions and warnings\n    -> Display results and record the decision')
table(['Component','Technology','Purpose'],[
('Frontend','React and Vite','Scenario inputs and four stage dashboard'),
('API','FastAPI and Pydantic','Validate requests and return analysis results'),
('Data processing','Pandas and NumPy','Read, join and prepare tabular data'),
('Prediction','Scikit learn and XGBoost','Train and compare severity classifiers'),
('Optimization','SciPy HiGHS','Solve bounded allocation linear programs'),
('Explainability','SHAP','Explain classifier outputs'),
('Map','Leaflet and OpenStreetMap','Display area locations and risk levels'),
('Alternative interface','Streamlit','Run the guided Python dashboard')],[1.35,1.8,3.5])
h('Source Organization')
p('The core directory contains data preparation, prediction, demand, priority, optimization, explanation, simulation and audit modules. api.py exposes the backend handlers. frontend/src contains the React application and styles. app/Home.py provides the Streamlit interface. The tests directory contains five automated test modules. Earlier modules remain in src and app/_pages_archive.')

page('Dataset Description')
purposes=[('damage.csv','Structural damage observations'),('population.csv','Population, coordinates and vulnerable age groups'),('disasters.csv','Event type, magnitude and event details'),('resources.csv','Available resource quantity by warehouse'),('warehouses.csv','Depot identifiers and locations'),('roads.csv','Area connections and road conditions'),('hospitals.csv','Hospital locations and capacity'),('shelters.csv','Shelter capacity and occupancy'),('scenarios.csv','Named example scenarios'),('historical_disaster_damage.csv','Labelled training examples')]
table(['File','Rows','Description'],[(name,inventory[name]['rows'],desc) for name,desc in purposes],[2.7,.5,3.45])
h('Feature Preparation')
p('Damage observations are joined with population records using area_id and disaster information using disaster_id. Missing damage counts are filled with zero. If building_damage is unavailable, it is reconstructed from minor, major and destroyed building counts. Missing population density and magnitude are filled with medians, with fixed fallback values when necessary.')
table(['Classifier input','Description'],[('building_damage','Total damaged buildings'),('minor_damage','Buildings with minor damage'),('major_damage','Buildings with major damage'),('destroyed_buildings','Buildings classified as destroyed'),('population_density','Population density of the area'),('magnitude','Magnitude recorded for the disaster')],[2.45,4.2])
p('The target classes are LOW, MEDIUM, HIGH and CRITICAL. The historical dataset contains 800 synthetic examples. The saved evaluation uses 640 records for training and 160 for the held out test set.')

page('Implementation and Setup')
h('Environment Setup')
p('Run the following commands from the emergency-resource-allocation directory. FastAPI and Uvicorn must be installed separately because they are not listed in the supplied requirements file.')
code('python -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npython -m pip install -r requirements.txt\npython -m pip install fastapi uvicorn')
h('Model Artifacts')
p('The project already includes best_model.pkl, damage_model.pkl, label_encoder.pkl, shap_explainer.pkl and metrics.json. The captured results use these artifacts. To train the three candidate models again, run the command below. Retraining replaces the model files and evaluation results.')
code('python core/model.py')
h('Run the Backend and Frontend')
p('Start the API from the project directory. In a second terminal, open the frontend directory and start Vite.')
code('python -m uvicorn api:app --host 127.0.0.1 --port 8000\n\n# Second terminal\ncd frontend\nnpm ci\nnpm run dev -- --host 127.0.0.1')
p('Open the local URL printed by Vite, normally http://127.0.0.1:5173. The React application calls the API at http://127.0.0.1:8000.')
h('Testing and the Streamlit Interface')
code('python -m pytest tests -v\nstreamlit run app/Home.py')
p('Streamlit normally opens at http://localhost:8501. The Python modules use paths relative to the project directory, so commands should be run from that directory.')

page('Processing Methodology')
for title,body in [
('Data Loading','load_all_data reads the configured CSV files. validate_dataset checks required columns and reports empty tables or negative values. prepare_ml_dataset merges the main tables and prepares the six classifier inputs.'),
('Severity Prediction','The training module compares Logistic Regression, Random Forest and XGBoost using a stratified 80 to 20 split and five fold cross validation. The implementation selects the model with the highest held out macro F1. During analysis, it loads the saved model and decodes predicted class labels.'),
('Demand Estimation','The predicted severity determines the affected population ratio. Resource demand is calculated from affected population, daily consumption rate and response duration. Duration supplied in hours is converted to days by the API.'),
('Priority Calculation','The priority module combines structural impact, demand, a shortage estimate, vulnerable population and road access obstruction. Component values and the final score are retained so that the result can be inspected.'),
('Resource Allocation','A linear program is solved for each resource. Total allocation is limited by available stock, while each area is capped at its own demand. A population proportional allocation is calculated as the comparison baseline.'),
('Explanation and Output','The system creates action cards and warnings from allocation results. The explanation module calculates local and global SHAP values and shows the components of a priority score. The API returns the results as JSON for the React dashboard.'),
('Simulation and Audit','The pipeline can be rerun after supply or population changes. The analysis handler writes a CSV record with a timestamp, scenario, model, solver details and allocation totals. Calling the core pipeline alone does not automatically write an audit record.')]:
    h(title);p(body)

page('Demand and Priority Calculations')
h('Affected Population and Demand')
code('affected = round(population * severity_ratio)\ndemand = affected * daily_resource_rate * duration_days\nshortage = max(0, demand - allocation)\ncoverage = min(1, total_supply / total_demand)')
table(['Severity','Affected population ratio'],[('LOW','0.10'),('MEDIUM','0.30'),('HIGH','0.55'),('CRITICAL','0.85')],[3.3,3.35])
p('The default core rates are 2 food rations, 3 litres of water and 0.05 medical kits per affected person per day, for three days. These rates and the severity ratios are project assumptions. The API uses different rates: food remains 2, water is 2 or 2.5 for Wildfire, and medical kits are 0.12, 0.18 for Wildfire or 0.25 for Earthquake.')
h('Priority Score')
code('priority = 0.25 * normalized_impact\n         + 0.25 * normalized_demand\n         + 0.20 * normalized_shortage\n         + 0.15 * normalized_medical_need\n         + 0.15 * access_penalty')
p('Impact combines building damage with a severity weight. Medical need is calculated as 1.5 times elderly population plus 1.2 times children population. If total_shortage is unavailable, the shortage input is half of estimated demand. The first four components use min max normalization. Road obstruction is already bounded between zero and one, using OPEN = 0, DAMAGED = 0.5 and BLOCKED = 1.')
p('Scores below 0.25 are LOW, scores from 0.25 to below 0.50 are MEDIUM, scores from 0.50 to below 0.75 are HIGH, and scores of 0.75 or more are CRITICAL. Operational priority is distinct from the severity class predicted by the classifier.')

page('Resource Allocation Model')
p('For one resource, x_i is the quantity allocated to area i, D_i is its demand, S is total stock, and p_i is the priority score with a minimum coefficient of 0.05. The allocation problem is:')
code('Maximize   sum_i (p_i / D_i) * x_i\n\nSubject to sum_i x_i <= S\n           0 <= x_i <= D_i')
p('SciPy linprog minimizes the negative of this objective using the HiGHS solver. The model rewards priority weighted demand satisfaction. Since benefit is divided by demand, an area with the highest priority score does not necessarily receive the first allocation. A positive objective coefficient does not guarantee a positive allocation.')
h('Baseline Allocation')
code('baseline_i = min(D_i, S * population_i / total_population)\nreduction = max(0, baseline_shortage - optimized_shortage)\nreduction_percent = 100 * reduction / baseline_shortage')
p('The comparison baseline distributes stock in proportion to resident population and caps the result at demand. In the default ten area experiment, both methods use the available food and water stock, and the reported total shortage reduction is zero. Allocation can change without reducing total shortage when the same scarce stock is fully consumed.')
h('Warehouse Allocation')
p('A separate transportation LP supports warehouse to area quantities and includes a small great circle distance penalty. It checks depot stock and area demand. The default core pipeline does not call this routine, and blocked road paths are not constraints in this transportation model.')
h('Units and Rounding')
p('Food is measured in rations, water in litres and medical supplies in kits. The core rounds quantities to one decimal place; the API rounds displayed area values to integers. The program also reports a combined shortage KPI that adds different resource units. The resource specific results below are used for interpretation.')

page('Results of the Core Analysis')
p('The baseline uses all ten areas, the supplied model artifacts, the original warehouse quantities and a three day response duration. No population or supply shock is applied.')
table(['Measure','Result'],[('Resident population','30,260'),('Estimated affected residents','12,846'),('Areas with CRITICAL priority','1'),('Areas with HIGH priority','1'),('Food solver','OPTIMAL'),('Water solver','OPTIMAL'),('Medical kit solver','OPTIMAL')],[3.75,2.9])
table(['Resource','Demand','Supply','Allocated','Shortage','Coverage'],[(r.replace('_',' '),num(base['total_demands'][r],1),num(base['total_supplies'][r],1),num(sum(float(x['opt_alloc_'+r]) for x in rows),1),num(sum(float(x['opt_shortage_'+r]) for x in rows),1),f"{base['coverages'][r]*100:.2f}%") for r in ['Food','Water','Medical_Kit']],[1.1,1.15,1.1,1.1,1.1,1.1])
p('Food and water stock are fully allocated, but they cover only a small part of estimated demand. Medical stock exceeds demand, leaving 3,073 kits unused. All three resource problems return OPTIMAL for this experiment. The reported shortage reduction against the population baseline is 0.00 percent.')
h('Example Calculation for Area A')
p('Area A has 2,368 residents and predicted MEDIUM severity. Its affected population is round(2368 × 0.30) = 710. Water demand is 710 × 3 × 3 = 6,390 litres. The optimizer allocates 6,390 litres, giving zero water shortage for Area A.')
p('The complete values are retained in baseline_summary.json and baseline_allocations.csv in the evidence folder.')

page('Area Allocation Results')
table(['Area','Predicted severity','Priority','Priority level','Affected'],[(r['area_id'],r['predicted_severity'],r['priority_score'],r['priority_level'],num(r['estimated_affected'])) for r in sorted(rows,key=lambda x:-float(x['priority_score']))],[.5,1.65,1.0,1.6,1.9])
table(['Area','Food rations','Water litres','Medical kits','Water shortage'],[(r['area_id'],num(r['opt_alloc_Food'],1),num(r['opt_alloc_Water'],1),num(r['opt_alloc_Medical_Kit'],1),num(r['opt_shortage_Water'],1)) for r in rows],[.5,1.4,1.45,1.45,1.85])
p('Area E has the highest priority score, 0.7858, but receives zero food and water in this run. This follows from the priority divided by demand objective and the absence of a minimum allocation constraint. The result shows why allocation policy must be reviewed alongside solver status.')
p('One generated recommendation sends 4,313 litres of water to Area C. Its water demand is 9,324 litres, leaving 5,011 litres unmet after allocation. Recommendation quantities come from the optimizer; they do not represent confirmed deliveries.')

page('Model Evaluation')
p('The following results are recorded in the supplied models/metrics.json. They are the saved training evaluation, rather than a new training run. Macro F1 is the selection metric.')
table(['Model','Accuracy','Macro F1','CV macro F1','ROC AUC'],[(name.replace('_',' '),f"{m['accuracy']*100:.2f}%",f"{m['f1_macro']:.4f}",f"{m['cv_5fold_f1_macro_mean']:.4f}",f"{m['roc_auc_ovr_macro']:.4f}") for name,m in metrics['models'].items()],[2.45,1.0,1.0,1.2,1.0])
p('XGBoost has the highest saved held out macro F1, 0.7366, and recorded accuracy of 81.87 percent. Random Forest has the highest mean cross validation macro F1, 0.6697. The current selection uses the held out result, so it should not be described as a selection based on cross validation.')
h('XGBoost Confusion Matrix')
p('Rows show actual classes and columns show predicted classes.')
table(['Actual','CRITICAL','HIGH','LOW','MEDIUM'],[(name,*vals) for name,vals in zip(metrics['metadata']['classes'],metrics['models']['XGBoost']['confusion_matrix'])],[1.85,1.2,1.2,1.2,1.2])
h('Local SHAP Output for Area A')
table(['Feature','Input','SHAP value'],[(c['feature'],num(c['feature_value'],1 if c['feature']=='magnitude' else 0),f"{c['shap_value']:+.4f}") for c in local['contributions']],[3.4,1.6,1.65])
p('Area A is predicted as MEDIUM. Positive SHAP values support this predicted class; they do not establish physical causes or an ordinal increase in disaster risk. Contributions describe model behaviour, not causation.')

page('Scenario Analysis')
p('Two separate changes were tested against the default core run: a 30 percent reduction in water supply and a 25 percent increase in population. Other settings were held at the core defaults for each experiment.')
table(['Resource','Baseline coverage','Water shock coverage'],[(r.replace('_',' '),f"{base['coverages'][r]*100:.2f}%",f"{shock['coverages'][r]*100:.2f}%") for r in ['Food','Water','Medical_Kit']],[2.25,2.2,2.2])
doc.add_picture(str(OUT/'coverage_comparison.png'),width=Inches(6.65))
p('Figure 1  Resource coverage before and after the water supply reduction','Caption')
table(['Water measure','Baseline','After shock'],[('Supply in litres','20,000','14,000'),('Demand in litres','115,614','115,614'),('Shortage in litres','95,614','101,614'),('Affected residents','12,846','12,846')],[3.25,1.7,1.7])
p(f"The water shock increases unmet water demand by 6,000 litres. In the separate population experiment, estimated affected population increases to {surge['kpis']['total_affected_population']:,}. Food coverage falls to {surge['coverages']['Food']*100:.2f} percent and water coverage to {surge['coverages']['Water']*100:.2f} percent.")

page('Dashboard for Disaster Inputs')
p('The Describe screen collects disaster type, severity, response duration, available supplies and population values for seven areas. The captured scenario uses Wildfire, High severity and 72 hours. The displayed area populations total 21,779.')
figure('01_describe.jpg','Figure 2  Disaster description and resource input screen')
h('Input Handling')
p('Population edits are sent to the analysis API, which updates the seven study areas before running the pipeline. Hours are converted into days. The UI displays food, water and medical stock inputs, but the current API writes these to total_quantity while the optimizer reads quantity. The displayed stock inputs therefore do not change the stock used by allocation.')
p('The population values in this screen differ from the Pydantic request defaults and from the original ten area CSV data. Dashboard figures should be interpreted using the scenario shown in the screenshot.')

page('Dashboard for AI Analysis')
p('The AI analysis screen combines summary values, the geographical map and the area ranking. The captured run shows 21,779 residents, no CRITICAL priority areas, four HIGH priority areas and 28 percent water coverage.')
figure('02_ai_analysis.jpg','Figure 3  AI analysis dashboard with the Houston risk map and area ranking')
h('Reading the Results')
p('Area E ranks first with a displayed priority score of 0.66, followed by Areas C and F at 0.65. Markers use the area priority colour and can be selected to inspect a location. The map uses Leaflet with OpenStreetMap tiles.')
p('The dashboard label TOTAL AFFECTED currently sums resident population. It is not the severity adjusted estimated_affected value. The stock reserve label repeats the interface input, while the coverage calculation uses the CSV stock quantity.')

page('Dashboard for the Response Plan')
p('The response screen turns allocations into recommendation cards. Each card shows the resource quantity, the reason for its priority and the remaining shortage after allocation.')
figure('03_response_plan.jpg','Figure 4  Response plan with the generated resource recommendation cards')
h('Example Recommendation')
p('For the captured React scenario, the first Area E card recommends 10,212 food rations against demand of 10,560. It reports 348 rations still unmet. The following cards show 13,200 litres of water and 950.4 medical kits for the same area.')
p('The 25 percent survival floor statement visible in the screen is not enforced by the current LP. The allocation lower bound is zero. The impact banner also uses the aggregate comparison returned by the backend, although its wording refers to critical and high zones. These screen statements are implementation gaps, not guarantees established by the results.')
p('AUTHORIZE DISPATCH changes the local interface state only. There is no connected convoy control or delivery confirmation service.')

page('Dashboard for Simulation and Audit History')
p('The What if screen tests supply and population changes and displays the decision history. A 30 percent water reduction was executed before capturing the screen.')
figure('04_what_if_and_audit.jpg','Figure 5  Water supply simulation result and decision audit history')
h('Simulation Output')
p('The API simulation reports water coverage falling from 27 percent to 19 percent, a decrease of eight displayed percentage points. This handler reloads seven areas with CSV populations and core default rates. It does not carry forward the Describe screen inputs, so it has a different baseline from the 28 percent coverage shown in the AI analysis screenshot.')
p('The audit table displays existing timestamped scenario records. The bulletin text and broadcast button are interface demonstrations; clicking the button does not transmit an alert to citizens.')

page('Testing and Verification')
table(['Check','Result'],[('Existing Python tests','24 passed with one pandas FutureWarning'),('React production build','Completed with Vite 8.3.2'),('Core allocation','All three resource solvers returned OPTIMAL'),('Water shock','Water shortage increased by 6000 litres'),('Population surge','Demand increased and the pipeline completed'),('Warehouse routine','OPTIMAL with 20000 litres shipped from W1'),('Dashboard workflow','Four screens opened and water stress test completed'),('API stock comparison','Inputs reduced to 1 left allocations unchanged')],[2.35,4.3])
h('Automated Test Coverage')
table(['Test module','Tests','Main checks'],[('test_actions_and_explain.py','5','Action quantities, reasons, SHAP and score components'),('test_demand.py','4','Demand arithmetic, bounds and zero population'),('test_optimizer.py','7','Stock and demand limits, solver output and transport'),('test_priority.py','4','Bounds, categories, determinism and normalization'),('test_simulate_and_audit.py','4','Supply shock, comparison and audit persistence')],[2.7,.5,3.45])
p('The warning comes from filling missing values in an object array in core/explain.py. It concerns future pandas behaviour and did not fail a test. The tests do not cover every interface statement or API input connection.')
h('Execution Environment')
p('Verification used Python '+runtime['python']+' on Windows 11. Main versions were Pandas '+runtime['packages']['pandas']+', NumPy '+runtime['packages']['numpy']+', SciPy '+runtime['packages']['scipy']+', XGBoost '+runtime['packages']['xgboost']+', SHAP '+runtime['packages']['shap']+' and FastAPI '+runtime['packages']['fastapi']+'.')
p('The production build transformed 1,938 modules. API handlers were also executed directly with Pydantic requests. The dashboard screenshots show a separate browser run of the React workflow.')

page('Limitations and Future Work')
for title,body in [
('Synthetic Data','Damage and population records are generated examples. The model has not been validated on independently collected operational disaster data.'),
('Stock Input Connection','The API must update the quantity field consumed by the optimizer. Requested totals also need a defined distribution across warehouse rows.'),
('Minimum Service Policy','The current allocation can leave an area with zero supply. Any minimum service policy must be written as a feasible LP constraint and tested under severe scarcity.'),
('Scenario Consistency','The simulation handler reloads a separate baseline. It should receive the current analysis inputs so its before and after results describe the same scenario.'),
('Status and Population Labels','The API defaults a missing singular solver_status field to OPTIMAL. It should return actual per resource statuses and is_optimal. The affected population label should be aligned with estimated_affected.'),
('Prediction Evaluation','The current model selection uses the held out test metric. A separate validation procedure would leave the final test set independent of selection.'),
('Audit and Routing','CSV logs can be edited or cleared, so they are not tamper proof. The transportation LP penalizes distance but does not enforce a road network route around blocked links.'),
('Operational Integrations','Dispatch and broadcast actions are local state changes. A deployed system would need confirmed resource movement, authenticated users and connected communication services.')]:
    h(title);p(body)

page('Conclusion and Source Reference')
p('ReliefPlan AI implements the main stages of an emergency resource planning workflow: data preparation, severity prediction, demand estimation, priority scoring, constrained allocation and scenario comparison. The dashboards provide a way to inspect the inputs, locate affected areas and review recommended quantities.')
p('The tests and captured runs show that the processing modules execute with the supplied synthetic data and saved model artifacts. The water supply experiment illustrates the effect of reduced stock on resource coverage. The area results also show that an optimal solution depends on the chosen objective and policy constraints; solver status alone does not establish an acceptable distribution of aid.')
p('The next development work should connect stock inputs correctly, keep simulation inputs consistent with the current analysis and add a clear minimum service policy. The interface labels and displayed guarantees should then be checked against the computed results. Validation on independent disaster records is needed before operational use.')
h('GitHub Repository')
link(p())
h('Files Included with the Report')
p('The accompanying submission package contains the complete Python and React source, the input CSV files, trained model artifacts, tests, captured output and the four original dashboard images. The code appendix contains selected functions without printed line numbers.')
h('Reproducing the Output')
p('After extracting the package, install the project dependencies and run the tests from emergency-resource-allocation. To reproduce the JSON and CSV output, run python evidence/capture_outputs.py from the extracted package root. New verification audit records are kept separately from the project audit history.')
h('Source Files Used for the Results')
p('Baseline allocations and supply shock results come from core/simulate.py, core/optimizer.py and the packaged evidence CSV and JSON files. Model evaluation values come from models/metrics.json. Dashboard screenshots were captured from the running React frontend connected to api.py.')

# Unnumbered, exact code excerpts. The complete modules are in the archive.
selections=[('Data Preparation','core/data.py','prepare_ml_dataset'),('Affected Population','core/demand.py','estimate_affected_population'),('Resource Demand','core/demand.py','compute_area_demands'),('Priority Calculation','core/priority.py','calculate_area_priorities'),('Resource Allocation','core/optimizer.py','optimize_resource_allocation'),('Damage Prediction','core/model.py','predict_damage_severity'),('Decision Audit','core/audit.py','log_decision_run')]
listings=[]
for title,path,name in selections:
    text=(ROOT/path).read_text(encoding='utf-8'); tree=ast.parse(text)
    node=next(x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name==name)
    listings.append((title,path,text.splitlines()[node.lineno-1:node.end_lineno]))
lines=(ROOT/'frontend/src/App.jsx').read_text(encoding='utf-8').splitlines()
start=next(i for i,line in enumerate(lines) if 'const runAnalysis = async' in line)
end=next(i for i in range(start+1,len(lines)) if lines[i]=='  };')
listings.append(('React Analysis Request','frontend/src/App.jsx',lines[start:end+1]))
for title,path,lines in listings:
    printed=[]
    for line in lines:
        parts=textwrap.wrap(line.expandtabs(4),width=96,replace_whitespace=False,drop_whitespace=False,break_long_words=False,break_on_hyphens=False) or ['']
        printed.extend(parts)
    # Divide near blank source lines where possible to keep related statements together.
    offset=0; chunk=0
    while offset<len(printed):
        limit=min(offset+57,len(printed))
        if limit<len(printed):
            blanks=[i for i in range(offset+44,limit+1) if not printed[i].strip()]
            if blanks: limit=blanks[-1]
        page(title+' Code'+(' Continued' if chunk else ''))
        p('Source file  '+path)
        for line in printed[offset:limit]: p(line,'Program code')
        offset=limit
        while offset<len(printed) and not printed[offset].strip(): offset+=1
        chunk+=1

dest=SUB/'ReliefPlan_AI_Final_Report.docx'
doc.save(dest)
(OUT/'final_report_structure.json').write_text(json.dumps({'pages':chapters,'dashboard_files':[x.name for x in SHOTS.glob('*.jpg')]},indent=2),encoding='utf-8')
print(dest)
print('Planned pages:',len(chapters))
