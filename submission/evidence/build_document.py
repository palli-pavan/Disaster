from pathlib import Path
import ast, csv, json, textwrap, zipfile, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
SUB = OUT.parent
ROOT = SUB.parent / 'emergency-resource-allocation'
def js(n): return json.loads((OUT/n).read_text(encoding='utf-8'))
base, shock, surge = [js(n+'_summary.json') for n in ['baseline','water_shock','population_surge']]
api = js('api_analysis.json')
api_sim = js('api_simulation.json')
metrics = json.loads((ROOT/'models/metrics.json').read_text())
inventory = js('dataset_inventory.json')
runtime = js('runtime.json')
rows = list(csv.DictReader((OUT/'baseline_allocations.csv').open(encoding='utf-8')))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.67)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.3)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Calibri'
    styles[name].font.color.rgb=RGBColor(0,0,0)
    for border in styles[name].element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.08
styles['Title'].font.size=Pt(29)
styles['Title'].paragraph_format.space_after=Pt(16)
styles['Subtitle'].font.size=Pt(14)
styles['Heading 1'].font.size=Pt(18)
styles['Heading 1'].paragraph_format.space_after=Pt(10)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].paragraph_format.space_before=Pt(12)
styles['Heading 2'].paragraph_format.space_after=Pt(6)
code_style=styles.add_style('Source code',1)
code_style.font.name='Consolas'
code_style.font.size=Pt(8)
code_style.paragraph_format.line_spacing=Pt(10.2)
code_style.paragraph_format.space_after=Pt(0)
code_style.paragraph_format.keep_with_next=False
code_style.paragraph_format.widow_control=False
code_style.paragraph_format.left_indent=Inches(.02)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('ReliefPlan AI  |  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='ReliefPlan AI Project Submission'
doc.core_properties.subject='Code implementation steps and verified output'
doc.core_properties.author=''

page_titles=[]
def p(t='',style=None): return doc.add_paragraph(t,style)
def h(t,l=2): return doc.add_heading(t,l)
def page(title):
    if page_titles: doc.add_page_break()
    page_titles.append(title)
    h(title,1)
def table(headers,data,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,txt in enumerate(headers): t.rows[0].cells[i].text=str(txt)
    for row in data:
        cells=t.add_row().cells
        for i,txt in enumerate(row): cells[i].text=str(txt)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trPr.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
        for ci,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'243746' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'85'); x.set(qn('w:type'),'dxa'); mar.append(x)
            pr.append(mar)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.02
                for run in para.runs:
                    run.font.size=Pt(9)
                    if ri==0: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    p()
    return t
def code(t):
    for line in t.strip('\n').splitlines():
        for wrapped in textwrap.wrap(line.expandtabs(4),width=98,replace_whitespace=False,drop_whitespace=False) or ['']:
            p(wrapped,'Source code')
    p()
def num(v,d=0): return f'{float(v):,.{d}f}'

def chart():
    im=Image.new('RGB',(1500,650),'white'); d=ImageDraw.Draw(im)
    f=lambda size: ImageFont.truetype('C:/Windows/Fonts/calibri.ttf',size)
    d.text((40,20),'Resource coverage under a water supply shock',font=f(39),fill='#111111')
    left,top,bottom,width=120,120,525,1250
    for tick in range(0,101,25):
        yy=bottom-(bottom-top)*tick/100
        d.line((left,yy,left+width,yy),fill='#D9DDE0',width=2)
        d.text((35,yy-15),f'{tick}%',font=f(26),fill='#444444')
    for i,r in enumerate(['Food','Water','Medical_Kit']):
        x=left+120+i*400
        for k,summary in enumerate([base,shock]):
            val=summary['coverages'][r]*100
            xx=x+k*110; yy=bottom-(bottom-top)*val/100
            d.rectangle((xx,yy,xx+90,bottom),fill=['#244B65','#C5793A'][k])
            d.text((xx-3,yy-35),f'{val:.2f}%',font=f(25),fill='#111111')
        d.text((x-5,550),{'Food':'Food','Water':'Water','Medical_Kit':'Medical kits'}[r],font=f(29),fill='#111111')
    for i,(label,color) in enumerate([('Baseline','#244B65'),('Water supply reduced by 30 percent','#C5793A')]):
        x=120+i*370; d.rectangle((x,610,x+22,632),fill=color); d.text((x+33,603),label,font=f(26),fill='#222222')
    im.save(OUT/'coverage_comparison.png')
chart()

# Main report
p('ReliefPlan AI Project Submission','Title')
p('Emergency Resource Allocation and Decision Support System','Subtitle')
p('Code  |  Output  |  Implementation steps')
p('Submission date  16 October 2026')
h('Project abstract')
p('ReliefPlan AI converts disaster damage surveys, population data and regional stockpiles into severity predictions, area priorities and resource allocation recommendations. The project combines a trained XGBoost classifier with a SciPy linear programming optimizer, then presents results through React and Streamlit interfaces. Human emergency managers review the recommendations and make final decisions.')
p('This submission explains the implemented data flow, setup procedure, mathematical rules and main source code. It also presents results from executing the supplied code and trained artifacts, including area allocations, resource coverage, supply shock simulation, model evaluation records and automated tests.')
h('Verified outcome')
p('The existing test suite passed all 24 tests. The React production build also completed. With the default three day core scenario, 12,846 residents are estimated to need assistance across ten areas. The optimizer allocated 15,000 food rations, 20,000 litres of water and 1,927 medical kits. Water coverage fell from 17.30 percent to 12.11 percent after a 30 percent water supply reduction.')
p('The current system is an educational prototype using synthetic data. Its measured output and implementation limits are explained in this report so that the submission can be evaluated accurately.')
h('Submission files')
p('ReliefPlan_AI_Project_Submission.docx contains the report and selected source listings. ReliefPlan_AI_Code_and_Output.zip contains the full project source, input data, trained models, execution evidence and reproduction instructions.')
page_titles.append('Title and abstract')

page('Report guide and project objectives')
p('The report is organized so an examiner can understand the project, reproduce its results and inspect the code behind the output.')
table(['Section','Coverage'],[
('Project design','Architecture, technology and file organization'),
('Input data','CSV inventory, features and training data'),
('Implementation','Environment setup and processing stages'),
('Mathematical rules','Affected population, demand, priority and optimization'),
('Program output','Baseline allocation, scenario comparisons and model results'),
('Verification','Tests, API handler execution and frontend build'),
('Limits and demonstration','Current gaps, examiner walkthrough and submission checklist'),
('Code appendix','Selected implementation listings with file and line references')],[1.6,5.2])
h('Problem statement')
p('Disaster response requires decisions about how limited food, water and medical supplies should be distributed among communities with different levels of damage and need. Manual allocation can obscure the assumptions behind these decisions and make scenario comparison difficult.')
h('Objectives')
for t in ['Predict damage severity from structural damage and disaster features.','Estimate affected population and resource demand for a chosen response duration.','Rank areas with a transparent weighted priority score.','Allocate resources within available stock and area demand limits.','Explain predictions and priorities, compare supply or population shocks, and record decisions.']:
    p(t,'List Bullet')
h('Operating scope')
p('The default core analysis uses all ten study areas A to J. The React API analysis filters the data to seven areas A to G. These two paths use different population inputs and consumption assumptions, so their output is reported separately.')

page('System architecture and source organization')
p('The application separates computation from presentation. Both interfaces call reusable functions in core. The React interface sends JSON requests to FastAPI, while Streamlit calls the pipeline from Python.')
code('CSV data and saved models\n          |\n          v\nLoad and validate -> Merge features -> Predict severity\n          |\n          v\nEstimate affected population -> Calculate demand -> Score priorities\n          |\n          v\nSolve allocation LP -> Build actions and warnings -> Summarize results\n          |\n          v\nReact through FastAPI or Streamlit -> Human review -> CSV audit history')
table(['Path','Implemented responsibility'],[
('api.py','FastAPI request models and health, analysis, simulation, metrics and audit handlers'),
('core/data.py and core/model.py','CSV validation, feature preparation, training and inference'),
('core/demand.py and core/priority.py','Affected population, resource demand and area scoring'),
('core/optimizer.py','Resource allocation LP, population baseline and transportation LP'),
('core/actions.py and core/explain.py','Action recommendations, warnings, SHAP and score decomposition'),
('core/simulate.py and core/audit.py','Full pipeline, scenario comparison and CSV decision records'),
('frontend/src and app/Home.py','React four step interface and Streamlit guided dashboard'),
('tests and models','Five test modules and saved model artifacts')],[2.7,4.1])
p('The src folder contains earlier modules. The app/_pages_archive folder contains earlier Streamlit pages; it is not an active app/pages directory. These files are retained in the source archive, while this report describes the current core pipeline and guided interfaces.')

page('Input datasets and machine learning features')
purposes={
'damage.csv':'Damage observations keyed by area_id and disaster_id',
'population.csv':'Coordinates, population, density and vulnerable age groups',
'disasters.csv':'Event type, date, magnitude and reported event impacts',
'resources.csv':'Resource type and quantity at each warehouse',
'warehouses.csv':'Depot identifiers, names and coordinates',
'roads.csv':'Connections, distance, travel time and road condition',
'hospitals.csv':'Hospital location, available beds and staffing',
'shelters.csv':'Shelter capacity and occupancy',
'scenarios.csv':'Named example scenarios',
'historical_disaster_damage.csv':'Synthetic labelled model training records'}
table(['CSV file','Rows','Role'],[(name,inventory[name]['rows'],purpose) for name,purpose in purposes.items()],[2.65,.5,3.65])
p('All supplied disaster impact and population records are synthetic. The coordinates describe a Houston study area. Hospital and shelter records are available as supporting data; the current run_full_pipeline function does not perform separate hospital bed or shelter assignment.')
h('Feature preparation')
p('Damage is joined to population on area_id and to disaster attributes on disaster_id. Missing damage counts become zero; building damage can be reconstructed from component counts. Missing density and magnitude use dataset medians, with fixed defaults when a median is unavailable.')
p('The six classifier inputs are building_damage, minor_damage, major_damage, destroyed_buildings, population_density and magnitude. The target is damage_severity with LOW, MEDIUM, HIGH and CRITICAL classes. Predicted severity and operational priority are separate fields.')

page('Implementation steps for setup and execution')
h('Step 1 Prepare the Python environment')
p('Open PowerShell in the emergency-resource-allocation directory. Create an isolated environment and install the project requirements. The React API also needs FastAPI and Uvicorn, which are absent from the supplied requirements file.')
code('cd D:\\Disaster\\emergency-resource-allocation\npython -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npython -m pip install -r requirements.txt\npython -m pip install fastapi uvicorn')
h('Step 2 Use the existing model artifacts')
p('The supplied models directory already contains best_model.pkl, damage_model.pkl, label_encoder.pkl, shap_explainer.pkl and metrics.json. These artifacts were used for the verification results in this report. To train again when needed, run the command below; retraining replaces the model and metric artifacts.')
code('python core/model.py')
h('Step 3 Run the tests')
code('python -m pytest tests -v')
h('Step 4 Launch the React application')
p('Run the API from the project directory so its relative data and model paths resolve correctly. Keep the API terminal open, then start Vite in a second terminal.')
code('python -m uvicorn api:app --host 127.0.0.1 --port 8000\n\n# In a second terminal\ncd frontend\nnpm ci\nnpm run dev -- --host 127.0.0.1')
p('Open the local URL printed by Vite, normally http://127.0.0.1:5173. The frontend sends requests to http://127.0.0.1:8000. A production build can be generated with npm run build.')
h('Step 5 Launch Streamlit when required')
code('streamlit run app/Home.py')
p('The Streamlit interface normally opens at http://localhost:8501. React and Streamlit provide alternate interfaces to the shared computation modules.')

page('Implementation steps for the processing pipeline')
steps=[
('Step 6 Load and validate data','load_all_data reads the configured CSV files and validate_dataset checks required columns. Negative numeric values and empty tables produce warnings. prepare_ml_dataset joins the tables and fills missing feature values.'),
('Step 7 Train or load the severity model','Training uses an 80 percent stratified training set and 20 percent held out set with random_state 42. Logistic Regression, Random Forest and XGBoost are evaluated. The implementation selects the model with the largest held out macro F1 and saves the model, label encoder, SHAP explainer and metrics.'),
('Step 8 Estimate demand','Load the saved model and decode predictions. Map each severity to an affected population ratio. Multiply affected residents by each resource rate and response duration. The API converts duration_hours to days before calling the pipeline.'),
('Step 9 Calculate area priority','Combine structural impact, estimated demand, an initial shortage proxy, demographic medical need and road access obstruction. Normalize the weights and classify the final score into a risk tier.'),
('Step 10 Solve resource allocation','Solve one bounded linear program per resource with SciPy HiGHS. Supply constrains the total allocation and demand constrains each area. Generate a population proportional baseline for comparison.'),
('Step 11 Explain and present results','Generate action cards and shortage or road obstruction warnings. The explain module supports global and local SHAP values and a component breakdown of priority scores. React displays maps, ranked areas, allocations and scenario results.'),
('Step 12 Compare scenarios and log decisions','The simulator reruns the same pipeline after modifying supply or population. The API analysis handler appends a CSV audit record with model, inputs, statuses and aggregate allocations. The core pipeline itself does not write an audit record automatically.')]
for title,body in steps: h(title); p(body)

page('Mathematical rules and configured assumptions')
h('Affected population and resource demand')
code('affected_i = round(population_i * severity_ratio_i)\ndemand_i_resource = affected_i * daily_rate_resource * duration_days\nshortage_i_resource = max(0, demand_i_resource - allocation_i_resource)\ncoverage_resource = min(1, total_supply_resource / total_demand_resource)')
table(['Severity','Affected population ratio'],[('LOW','0.10'),('MEDIUM','0.30'),('HIGH','0.55'),('CRITICAL','0.85')],[3.1,3.7])
table(['Parameter','Core default','API analysis default'],[('Duration','3 days','72 hours converted to 3 days'),('Food','2 rations per person per day','2 rations per person per day'),('Water','3 litres per person per day','2 litres or 2.5 for Wildfire'),('Medical kits','0.05 per person per day','0.12 or 0.18 for Wildfire or 0.25 for Earthquake')],[1.2,2.5,3.1])
h('Priority score')
code('priority = 0.25 * normalized_impact + 0.25 * normalized_demand\n         + 0.20 * normalized_shortage + 0.15 * normalized_medical_need\n         + 0.15 * access_penalty')
p('Impact uses building damage and a severity weight. Medical need uses 1.5 times elderly population plus 1.2 times children population. Until total_shortage exists, the shortage component is a proxy equal to half of estimated demand. The first four components use scenario min max scaling; access is already bounded from road status scores OPEN 0, DAMAGED 0.5 and BLOCKED 1.')
p('The priority tiers are LOW below 0.25, MEDIUM from 0.25 to below 0.50, HIGH from 0.50 to below 0.75, and CRITICAL from 0.75. Consumption rates, severity ratios and policy weights are project assumptions.')

page('Optimization formulation and result interpretation')
p('For each resource, the decision variable x_i is the quantity allocated to area i. D_i is demand, S is available supply and p_i is max(priority_score_i, 0.05). The implementation minimizes a negative benefit to express a maximization problem through linprog.')
code('Maximize   sum_i (p_i / D_i) * x_i\nSubject to sum_i x_i <= S\n           0 <= x_i <= D_i\n\nPython objective coefficient  c_i = -(p_i / D_i)\nPython solver                 linprog(..., method="highs")')
p('An optimal result maximizes priority weighted demand satisfaction. Because the marginal benefit is priority divided by demand, the area with the largest raw priority score is not necessarily served first. A positive priority weight does not guarantee a positive allocation.')
h('Baseline comparison')
code('baseline_allocation_i = min(D_i, S * population_i / total_population)\nbaseline_shortage_i = max(0, D_i - baseline_allocation_i)\nreduction = max(0, total_baseline_shortage - total_optimized_shortage)\nreduction_percent = 100 * reduction / total_baseline_shortage')
p('The default ten area run reports 0.00 percent total shortage reduction against this baseline. With scarce supply, both methods can consume the same available stock while assigning it differently. The seven area API example reports a small aggregate reduction, as shown in its separate output section. These observations do not support a general claim that optimization always reduces total shortage.')
h('Transportation extension')
p('optimize_transportation_lp adds warehouse to area quantities, stock constraints per depot and demand constraints per area. It subtracts a small great circle distance cost from demand satisfaction benefit. It is a separate callable routine and is not invoked by run_full_pipeline. It does not enforce blocked road path constraints.')
h('Units and rounding')
p('Food is reported in rations, water in litres and medical supplies in kits. The core rounds quantities to one decimal place, and the API rounds area values to integers. Combined shortage totals sum unlike units for a program KPI and should not be interpreted as a physical quantity or used to compare different resource types. Resource specific tables are more informative.')

page('Output from the default core scenario')
p('Execution inputs were all ten areas A to J, the saved classifier and encoder, the original CSV stockpiles, default priority parameters, three days of demand and no scenario shocks. These results were captured on 16 October 2026.')
table(['Measure','Observed value'],[
('Areas analyzed','10'),('Total resident population',num(sum(float(r['population']) for r in rows))),('Estimated affected population','12,846'),('Critical priority areas','1'),('High priority areas','1'),('Food solver status','OPTIMAL'),('Water solver status','OPTIMAL'),('Medical kit solver status','OPTIMAL'),('Baseline shortage reduction','0.00 percent')],[3.9,2.9])
table(['Resource','Demand','Supply','Allocated','Shortage','Coverage'],[
(r,num(base['total_demands'][r],1),num(base['total_supplies'][r],1),num(sum(float(x['opt_alloc_'+r]) for x in rows),1),num(sum(float(x['opt_shortage_'+r]) for x in rows),1),f"{base['coverages'][r]*100:.2f}%") for r in ['Food','Water','Medical_Kit']],[1.15,1.15,1.1,1.1,1.15,1.15])
p('The medical stockpile exceeds demand, so the optimizer leaves 3,073 kits unused. Food and water stock are fully used and large deficits remain. The solver returned OPTIMAL for each resource in this captured scenario.')
h('Hand checked example for Area A')
p('Area A has 2,368 residents and predicted MEDIUM severity. The affected ratio is 0.30, giving round(2368 × 0.30) = 710 residents. Water demand is 710 × 3 × 3 = 6,390 litres. The allocation is 6,390 litres, leaving zero water shortage for this area.')
p('Evidence files are baseline_summary.json and baseline_allocations.csv in the packaged evidence directory. These retain the unrounded values used to prepare the report tables.')

page('Output by area and recommended actions')
table(['Area','Predicted severity','Priority','Risk tier','Affected'],[(r['area_id'],r['predicted_severity'],r['priority_score'],r['priority_level'],num(r['estimated_affected'])) for r in sorted(rows,key=lambda r:-float(r['priority_score']))],[.5,1.7,1,1.4,2.3])
table(['Area','Food rations','Water litres','Medical kits','Water shortage'],[(r['area_id'],num(r['opt_alloc_Food'],1),num(r['opt_alloc_Water'],1),num(r['opt_alloc_Medical_Kit'],1),num(r['opt_shortage_Water'],1)) for r in rows],[.5,1.45,1.45,1.45,2.05])
p('Area E has the highest priority score at 0.7858, but receives zero food and water in this run. The LP objective rewards priority divided by demand and has no minimum service constraint. This is a consequential behavior to discuss during evaluation.')
h('Example generated recommendation')
card=next(c for c in base['action_cards'] if c.get('area_id')=='C' and c.get('resource')=='Water')
p(f"Dispatch {card['quantity']:,.1f} litres of water to Area C. {card['reason']} {card['expected_effect']}")
p('An action card is a recommendation generated from allocation output. It does not send a convoy or verify delivery.')

page('Output from supply and population simulations')
p('The core simulator was rerun with a 30 percent reduction in water stock and, independently, with a 25 percent population increase. Other baseline settings were retained for each experiment.')
table(['Resource','Baseline coverage','Water shock coverage'],[(r,f"{base['coverages'][r]*100:.2f}%",f"{shock['coverages'][r]*100:.2f}%") for r in ['Food','Water','Medical_Kit']],[2.3,2.25,2.25])
doc.add_picture(str(OUT/'coverage_comparison.png'),width=Inches(6.8))
p('Figure 1 Resource coverage before and after the water supply shock','Caption')
table(['Measure','Baseline','Water supply reduced'],[
('Water stock in litres',num(base['total_supplies']['Water']),num(shock['total_supplies']['Water'])),
('Water demand in litres',num(base['total_demands']['Water']),num(shock['total_demands']['Water'])),
('Water shortage in litres',num(sum(float(r['opt_shortage_Water']) for r in rows)),num(shock['total_demands']['Water']-shock['total_supplies']['Water'])),
('Estimated affected population',num(base['kpis']['total_affected_population']),num(shock['kpis']['total_affected_population']))],[3.2,1.8,1.8])
p(f"The water shock increases unmet water demand by 6,000 litres. In the separate population surge scenario, estimated affected population rises to {surge['kpis']['total_affected_population']:,}; food coverage becomes {surge['coverages']['Food']*100:.2f} percent and water coverage becomes {surge['coverages']['Water']*100:.2f} percent.")
p('Evidence files include water_shock_summary.json, population_surge_summary.json and the corresponding allocation CSVs.')

page('Model evaluation output and explainability')
p('The existing models/metrics.json records 800 synthetic training examples, a 640 record training partition and a 160 record held out partition. The saved selection metric is macro F1. This table reproduces the persisted evaluation record; the model was not retrained for this submission.')
table(['Model','Accuracy','Macro F1','CV macro F1','ROC AUC'],[(name.replace('_',' '),f"{m['accuracy']*100:.2f}%",f"{m['f1_macro']:.4f}",f"{m['cv_5fold_f1_macro_mean']:.4f}",f"{m['roc_auc_ovr_macro']:.4f}") for name,m in metrics['models'].items()],[2.5,1,1,1.3,1])
p('XGBoost is selected with held out macro F1 of 0.7366 and recorded accuracy of 81.87 percent. Random Forest has the highest mean cross validation macro F1 of 0.6697. The implementation selects using the held out metric, so cross validation performance and selection criterion should be distinguished. Accuracy is reproduced at the precision stored in the JSON file.')
h('XGBoost confusion matrix')
p('Rows are actual classes and columns are predicted classes. The class order is CRITICAL, HIGH, LOW and MEDIUM.')
table(['Actual class','CRITICAL','HIGH','LOW','MEDIUM'],[(name,*vals) for name,vals in zip(metrics['metadata']['classes'],metrics['models']['XGBoost']['confusion_matrix'])],[2,1.2,1.2,1.2,1.2])
h('Executed local SHAP explanation for Area A')
local=js('local_explanation.json')
table(['Feature','Input value','SHAP value'],[(c['feature'],num(c['feature_value'],1 if c['feature']=='magnitude' else 0),f"{c['shap_value']:+.4f}") for c in local['contributions']],[3.6,1.6,1.6])
p('These are contributions to the predicted MEDIUM class score. Positive values support that class and do not establish a physical cause or an ordinal increase in disaster risk. Contributions describe model behaviour, not causation.')

page('API analysis output and interface workflow')
p('The analysis request uses the API defaults: Wildfire, High, 72 hours, seven areas A to G and resident populations totaling 22,455. Rates are 2 food rations, 2.5 litres of water and 0.18 medical kits per affected resident per day. Handler execution used an isolated audit file for the new verification records.')
code(json.dumps(api['overview'],indent=2))
p('The total_affected value above is actually the sum of resident populations, not the severity adjusted estimated_affected field. Coverage is calculated from the CSV quantity column. The inputs nominally request 40,000 food, 40,000 water and 2,500 medical kits, but those fields currently write total_quantity rather than quantity and therefore do not change allocation stock.')
h('API simulation output')
code(json.dumps({k:v for k,v in api_sim.items() if k!='warning'},indent=2))
p('This simulation independently reloads seven areas with CSV populations and core default rates. It does not reuse the analysis screen populations, stock inputs, duration or disaster specific rates. Its rounded 27 to 19 percent water coverage is therefore different from the ten area core experiment.')
h('Four step user workflow')
p('Describe captures disaster details and population inputs. AI analysis displays map markers, priorities and allocations. Response plan shows recommendation cards. What if runs the shock request and displays a bulletin and audit history. Dispatch and bulletin activation are client state demonstrations; external dispatch and messaging services are not connected.')

page('Verification results and execution environment')
table(['Check','Observed result'],[
('Existing Python tests','24 passed with 1 pandas FutureWarning'),
('React production build','Vite 8.3.2 completed and transformed 1938 modules'),
('Default core pipeline','All three resource solvers returned OPTIMAL'),
('Water supply shock','Coverage decreased and water shortage increased by 6000 litres'),
('Population surge','Pipeline completed and demand increased'),
('Transportation routine','OPTIMAL with 20000 litres shipped from W1'),
('API handler execution','Health, analysis, simulation, metrics and audit handlers returned'),
('Low stock input comparison','Changing API stock inputs to 1 left allocations unchanged')],[2.25,4.55])
p('The API checks invoked handlers directly with Pydantic request objects. They verify computation and response construction, rather than HTTP transport or browser interaction. The React check is a production build, rather than an end to end UI test. These checks establish the scope of the evidence.')
h('Test suite coverage')
table(['Test module','Tests','Coverage'],[
('test_actions_and_explain.py','5','Action quantities, numeric reasons, local SHAP and priority decomposition'),
('test_demand.py','4','Hand checked arithmetic, bounds and zero population'),
('test_optimizer.py','7','Supply and demand limits, solver status, zero supply and transportation'),
('test_priority.py','4','Score bounds, categories, determinism and constant normalization'),
('test_simulate_and_audit.py','4','Water shock, reduction formula, scenario comparison and audit persistence')],[2.8,.5,3.5])
p('The warning concerns future pandas behavior when filling missing values in object arrays in core/explain.py. No test failed. Tests do not cover every UI claim or API wiring behavior.')
h('Verified runtime')
p('Python '+runtime['python']+' on Windows 11. Packages: '+', '.join(f'{n} {v}' for n,v in runtime['packages'].items())+'.')
p('Build artifacts were dist/index.html, a 26.60 kB CSS bundle and a 403.65 kB JavaScript bundle. Dependency installation is required after extracting the source archive because node_modules is not packaged.')

page('Current limitations and improvement steps')
limits=[
('Connect stock inputs to the optimizer','In api.py, update the quantity field used by core/optimizer.py and distribute requested total stock across warehouse rows without duplicating it. The current total_quantity assignment leaves original CSV supply unchanged.'),
('Implement a feasible minimum allocation policy','The response screen claims a 25 percent survival floor, but LP bounds currently start at zero. Decide how to ration when supply is below the desired floor and encode the chosen policy as constraints before making this claim.'),
('Carry analysis state into simulations','The simulation handler reloads original data and default assumptions. Pass the current populations, duration and stock configuration to make the before and after comparison match the current analysis.'),
('Connect scenario controls to the intended computation','The severity string is used in the analysis audit name, while predicted severity comes from the model. Disaster type changes consumption assumptions but does not replace the damage survey or event magnitude.'),
('Use explicit status and population fields','The API reads a missing singular solver_status key and defaults to OPTIMAL. Return actual solver_statuses and is_optimal. Rename total_affected or compute it from estimated_affected so the KPI matches its label.'),
('Separate prototype actions from operational execution','Dispatch and broadcast controls change React state only. The displayed convoy ETA is a fixed interface value. Real integrations would need delivery confirmation and event records.'),
('Strengthen reproducibility and evaluation','Pin dependencies, select models using training validation rather than the held out test set, and validate with independently collected data. The current training and evaluation records are synthetic.'),
('Strengthen audit and route guarantees','CSV append logging is editable and clear_audit_log can delete it; it is not tamper proof. The transportation LP uses distance penalties, while blocked road information currently produces warnings rather than route feasibility constraints.')]
for title,body in limits: h(title);p(body)

page('Demonstration and submission checklist')
h('Suggested examiner demonstration')
for title,body in [
('1 Introduce the problem','Explain limited stock, different area needs and the role of a human decision maker. Show the source folders and describe the synthetic data.'),
('2 Describe the scenario','Open React and show the seven population fields, duration and disaster type. Explain the current stock input limitation before interpreting coverage.'),
('3 Run and inspect analysis','Run analysis, select a map area and compare predicted severity with operational priority. Show quantities and remaining shortage.'),
('4 Explain the allocation','Use the objective and area table in this document to explain priority weighted demand satisfaction. Show that zero allocations are possible.'),
('5 Stress test water supply','Run a 30 percent water supply shock. Explain that the API simulation uses its independent default baseline; use the packaged core comparison for a controlled ten area experiment.'),
('6 Present validation','Show the 24 passing tests, saved model comparison and SHAP output. End with the current limitations and improvement steps.')]:
    h(title);p(body)
h('Files to submit')
for t in ['This Word report with implementation explanation, code listings and output tables.','The complete source and output ZIP, including data, models, tests and frontend source.','The evidence folder within the ZIP, including JSON output, area CSVs, execution notes and reproduction scripts.']:
    p(t,'List Bullet')
h('Reproduction of captured evidence')
p('After extracting the ZIP, install the dependencies and run python evidence/capture_outputs.py from the archive root. The script resolves the sibling emergency-resource-allocation directory and writes new verification output to evidence. Its new analysis records use evidence/verification_audit.csv, preserving the project audit history.')
p('The source appendix below contains selected complete Python functions and the React analysis request handler. File paths and source line ranges refer to the submitted project snapshot. All imports, remaining functions, interfaces, style sheets and tests are present in the source archive.')

# Exact source excerpts, wrapped only for print readability.
selected=[
('Data preparation','core/data.py','prepare_ml_dataset'),
('Affected population estimation','core/demand.py','estimate_affected_population'),
('Resource demand calculation','core/demand.py','compute_area_demands'),
('Area priority calculation','core/priority.py','calculate_area_priorities'),
('Linear programming allocation','core/optimizer.py','optimize_resource_allocation'),
('Severity prediction','core/model.py','predict_damage_severity'),
('Decision audit record','core/audit.py','log_decision_run')]
listings=[]
for title,path,func in selected:
    txt=(ROOT/path).read_text(encoding='utf-8')
    tree=ast.parse(txt)
    node=next(x for x in tree.body if isinstance(x,(ast.FunctionDef,ast.AsyncFunctionDef)) and x.name==func)
    listings.append((title,path,node.lineno,node.end_lineno,txt.splitlines()[node.lineno-1:node.end_lineno]))
txt=(ROOT/'frontend/src/App.jsx').read_text(encoding='utf-8').splitlines()
start=next(i for i,l in enumerate(txt) if 'const runAnalysis = async' in l)
end=next(i for i in range(start+1,len(txt)) if txt[i]=='  };')
listings.append(('React analysis request','frontend/src/App.jsx',start+1,end+1,txt[start:end+1]))
for idx,(title,path,start,end,lines) in enumerate(listings,1):
    wrapped=[]
    for n,line in enumerate(lines,start):
        parts=textwrap.wrap(line.expandtabs(4),width=91,replace_whitespace=False,drop_whitespace=False) or ['']
        wrapped.extend([f'{n:>3}  '+parts[0]]+['     '+q for q in parts[1:]])
    for chunk_idx,offset in enumerate(range(0,len(wrapped),57)):
        page(f'Code listing {idx} '+title+(' continued' if chunk_idx else ''))
        p(f'Source  {path}  |  Lines {start} to {end}')
        if chunk_idx==0:
            p('Exact source function with printed line numbers. Dependencies and surrounding module code are in the source archive.')
        for line in wrapped[offset:offset+57]: p(line,'Source code')

path=SUB/'ReliefPlan_AI_Project_Submission.docx'
doc.save(path)
(OUT/'document_structure.json').write_text(json.dumps({'planned_sections':page_titles,'code_listings':[(p,s,e) for _,p,s,e,_ in listings]},indent=2),encoding='utf-8')
print(str(path))
print('Planned pages:',len(page_titles))
