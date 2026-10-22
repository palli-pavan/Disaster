from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parent.parent
source=ROOT/'ReliefPlan_AI_Final_Report.docx'
doc=Document(source)
members=[
    ('PALLI PAVAN KUMAR','23MID0181'),
    ('DUDDU SATYA SAI VISHNU VARDHAN','23MID0163'),
    ('DARISI RITHVIK','23MID0209'),
    ('G TARUNI','23MID0170'),
]
anchor=next(p for p in doc.paragraphs if p.text=='16 October 2026')
texts=['TEAM MEMBERS']+[f'{i}) {name} – {reg}' for i,(name,reg) in enumerate(members,1)]
for i,text in enumerate(texts):
    element=OxmlElement('w:p')
    anchor._p.addnext(element)
    paragraph=Paragraph(element,anchor._parent)
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before=Pt(7 if i==0 else 0)
    paragraph.paragraph_format.space_after=Pt(5)
    paragraph.paragraph_format.line_spacing=1.0
    paragraph.paragraph_format.keep_with_next=i<4
    run=paragraph.add_run(text)
    run.font.name='Times New Roman'
    run.font.size=Pt(12 if i==0 else 11)
    run.font.bold=i==0
    run.font.color.rgb=RGBColor(0,0,0)
    anchor=paragraph

for style in doc.styles:
    if style.name.startswith('Heading') or style.name in ['Title','Subtitle']:
        style.font.bold=True
        style.font.color.rgb=RGBColor(0,0,0)
        if style.name=='Heading 2': style.font.size=Pt(13)
        for color in style.element.xpath('.//w:color'):
            for attr in ['themeColor','themeTint','themeShade']:
                color.attrib.pop(qn('w:'+attr),None)
            color.set(qn('w:val'),'000000')

for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith('Heading') or paragraph.style.name in ['Title','Subtitle']:
        for run in paragraph.runs:
            run.font.bold=True
            run.font.color.rgb=RGBColor(0,0,0)
            if paragraph.style.name=='Heading 2': run.font.size=Pt(13)
            for color in run._r.xpath('.//w:color'):
                for attr in ['themeColor','themeTint','themeShade']:
                    color.attrib.pop(qn('w:'+attr),None)
                color.set(qn('w:val'),'000000')

dest=ROOT/'ReliefPlan_AI_Team_Report.docx'
doc.save(dest)
checked=Document(dest)
text='\n'.join(p.text for p in checked.paragraphs)
assert all(name in text and reg in text for name,reg in members)
assert len(checked.inline_shapes)==5
assert all(r.font.bold and str(r.font.color.rgb)=='000000' for p in checked.paragraphs if p.style.name.startswith('Heading') for r in p.runs)
print(dest)
